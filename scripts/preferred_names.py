#!/usr/bin/env python3
"""Apply preferred names and countries from IB2 list (1).xlsx onto students.json."""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ROSTER_XLSX = ROOT / "IB2 list (1).xlsx"
HOUSES_XLSX = ROOT / "OPEN HOUSE DISTRIBUITION.xlsx"
STUDENTS_JSON = ROOT / "src" / "data" / "students.json"
EDGE_ROSTER = ROOT / "supabase" / "functions" / "_shared" / "roster.ts"
EMAILS_TS = ROOT / "src" / "data" / "studentEmails.ts"

HOUSE_NAMES = {
    "cabo blanco": "Cabo Blanco",
    "cahuita": "Cahuita",
    "coco": "Coco",
    "flamingo": "Flamingo",
    "hermosa": "Hermosa",
    "malpais": "Malpaís",
    "montezuma": "Montezuma",
    "tortuguero": "Tortuguero",
}

SKIP_PARTS = {"", "x", "-", "none", "n/a", "null"}
PARTICLES = {
    "a",
    "al",
    "bin",
    "binti",
    "da",
    "das",
    "de",
    "del",
    "della",
    "der",
    "di",
    "do",
    "dos",
    "du",
    "la",
    "las",
    "le",
    "los",
    "of",
    "the",
    "van",
    "von",
    "y",
}

DEMONYMS = {
    "afghan": "Afghanistan",
    "american": "United States",
    "belgian": "Belgium",
    "bermudian": "Bermuda",
    "bhutanese": "Bhutan",
    "brazilian": "Brazil",
    "british": "United Kingdom",
    "bulgarian": "Bulgaria",
    "burmese": "Myanmar",
    "canadian": "Canada",
    "chilean": "Chile",
    "chinese": "China",
    "colombian": "Colombia",
    "congolese": "Democratic Republic of the Congo",
    "costa rican": "Costa Rica",
    "czech": "Czech Republic",
    "danish": "Denmark",
    "dutch": "Netherlands",
    "french": "France",
    "german": "Germany",
    "greek": "Greece",
    "guatemalan": "Guatemala",
    "honduran": "Honduras",
    "hungarian": "Hungary",
    "indian": "India",
    "israeli": "Israel",
    "italian": "Italy",
    "jamaican": "Jamaica",
    "kenyan": "Kenya",
    "lithuanian": "Lithuania",
    "malaysian": "Malaysia",
    "mexican": "Mexico",
    "nepalese": "Nepal",
    "nicaraguan": "Nicaragua",
    "nigerian": "Nigeria",
    "norwegian": "Norway",
    "paraguayan": "Paraguay",
    "peruvian": "Peru",
    "polish": "Poland",
    "portuguese": "Portugal",
    "romanian": "Romania",
    "russian": "Russia",
    "rwandan": "Rwanda",
    "salvadoran": "El Salvador",
    "serbian": "Serbia",
    "sierra leonean": "Sierra Leone",
    "slovak": "Slovakia",
    "spanish": "Spain",
    "swiss": "Switzerland",
    "taiwanese": "Taiwan",
    "ugandan": "Uganda",
    "uruguayan": "Uruguay",
    "venezuelan": "Venezuela",
    "vietnamese": "Vietnam",
}

PAREN_COUNTRY = {
    "congo, the democratic republic of the": "Democratic Republic of the Congo",
    "united states": "United States",
    "netherlands": "Netherlands",
    "nigeria": "Nigeria",
    "norway": "Norway",
    "china": "China",
    "france": "France",
    "italy": "Italy",
    "india": "India",
    "myanmar": "Myanmar",
}


def strip_accents(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def normalize_name(name: str) -> str:
    name = name.replace("\t", " ").replace("\n", " ")
    name = strip_accents(name)
    name = name.replace("’", "'").replace("`", "'")
    name = re.sub(r"[^A-Za-z0-9'\- ]+", " ", name)
    name = re.sub(r"\s+", " ", name).strip().lower()
    return name


def tokens(name: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", normalize_name(name))


def token_set(name: str) -> set[str]:
    return set(tokens(name))


def clean_part(value: object) -> str:
    text = "" if value is None else str(value).strip()
    text = text.strip("().")
    text = re.sub(r"\s+", " ", text).strip()
    if text.lower() in SKIP_PARTS:
        return ""
    return text


def smart_case_word(word: str, index: int) -> str:
    lower = word.lower()
    if "'" in word and not word.isupper():
        return word
    if lower.replace("'", "") in PARTICLES and index > 0:
        if word.startswith("d'") or word.startswith("D'"):
            rest = word[2:]
            return "d'" + (rest[:1].upper() + rest[1:] if rest else "")
        return lower
    if "-" in word:
        return "-".join(
            (part[:1].upper() + part[1:].lower()) if part else part
            for part in word.split("-")
        )
    letters = [ch for ch in word if ch.isalpha()]
    if letters and (all(ch.isupper() for ch in letters) or all(ch.islower() for ch in letters)):
        return word[:1].upper() + word[1:].lower()
    return word


def smart_case(text: str) -> str:
    text = re.sub(r"-{2,}", "-", text)
    words = text.split()
    return " ".join(smart_case_word(word, i) for i, word in enumerate(words))


def accent_count(text: str) -> int:
    return sum(1 for ch in text if ord(ch) > 127)


def keep_richer_spelling(formatted: str, current: str) -> str:
    if tokens(formatted) != tokens(strip_parens(current)):
        return formatted
    current_core = strip_parens(current)
    if accent_count(formatted) >= accent_count(current_core):
        return formatted
    return current_core


def strip_parens(name: str) -> str:
    return re.sub(r"\s*\([^)]*\)", "", name).strip()


def country_from_nationality(raw: object) -> str | None:
    text = clean_part(raw)
    if not text:
        return None
    match = re.search(r"\(([^)]+)\)", str(raw))
    if match:
        inner = re.sub(r"\s+", " ", match.group(1)).strip()
        mapped = PAREN_COUNTRY.get(inner.lower())
        if mapped:
            return mapped
        return smart_case(inner)
    key = re.sub(r"\s+", " ", text).strip().lower()
    key = re.sub(r"\s*\(.*\)$", "", key).strip()
    return DEMONYMS.get(key, smart_case(text))


def last_spelling(last: str, current: str) -> str:
    if not last:
        return last
    cur = strip_parens(current)
    last_toks = tokens(last)
    cur_words = cur.split()
    n = len(last_toks)
    if n == 0 or len(cur_words) < n:
        return last
    flat: list[str] = []
    for word in cur_words:
        word_tokens = tokens(word)
        flat.append(word_tokens[0] if word_tokens else "")
    if flat[-n:] == last_toks:
        taken = " ".join(cur_words[-n:])
        letters = [ch for ch in taken if ch.isalpha()]
        if letters and all(ch.isupper() for ch in letters):
            return smart_case(taken)
        return taken
    return last


def infer_last(current: str, first: str, middle: str, preferred: str) -> str:
    cur_words = strip_parens(current).split()
    skip = token_set(" ".join(part for part in (first, middle, preferred) if part))
    leftover: list[str] = []
    for word in cur_words:
        word_tokens = token_set(word)
        if word_tokens and word_tokens.issubset(skip):
            continue
        leftover.append(word)
    return " ".join(leftover)


def leftover_given(first: str, middle: str, preferred: str) -> str:
    pref = token_set(preferred)
    leftover: list[str] = []
    for part in (first, middle):
        if not part:
            continue
        words: list[str] = []
        for word in part.split():
            word_tokens = token_set(word)
            if not word_tokens or word_tokens.issubset(pref):
                continue
            if word.lower() in PARTICLES and not words:
                continue
            words.append(word)
        while words and words[0].lower() in PARTICLES:
            words.pop(0)
        while words and words[-1].lower() in PARTICLES:
            words.pop()
        if words:
            leftover.append(" ".join(words))
    return " ".join(leftover)


def preferred_is_first(preferred: str, first: str) -> bool:
    first_word = first.split()[0] if first.strip() else ""
    return bool(preferred) and bool(first_word) and tokens(preferred) == tokens(first_word)


def include_middle(middle: str, current: str) -> bool:
    if not middle:
        return False
    words = [word for word in middle.split() if word.lower() not in PARTICLES]
    if len(words) <= 2:
        return True
    return token_set(middle).issubset(token_set(strip_parens(current)))


def format_display(
    preferred: str,
    first: str,
    middle: str,
    last: str,
    current: str,
) -> str:
    preferred = clean_part(preferred)
    first = clean_part(first)
    middle = clean_part(middle)
    last = clean_part(last)

    if last:
        last = last_spelling(smart_case(last) if last.isupper() else last, current)
        if last.isupper() or last.islower():
            last = last_spelling(smart_case(last), current)
    elif current:
        last = infer_last(current, first, middle, preferred)

    if first and (first.isupper() or first.islower()):
        first = smart_case(first)
    if middle and (middle.isupper() or middle.islower()):
        middle = smart_case(middle)
    if preferred and (preferred.isupper() or preferred.islower()):
        preferred = smart_case(preferred)

    legal_parts = [part for part in (first, middle, last) if part]
    legal = " ".join(legal_parts)
    if not preferred:
        return current or legal

    pref_tokens = token_set(preferred)
    legal_tokens = token_set(legal)
    if pref_tokens and pref_tokens == legal_tokens:
        return keep_richer_spelling(legal or preferred, current)
    if last and pref_tokens == token_set(f"{first} {last}"):
        return keep_richer_spelling(
            " ".join(part for part in (first, last) if part),
            current,
        )

    if preferred_is_first(preferred, first):
        parts = [first]
        if include_middle(middle, current):
            parts.append(middle)
        if last:
            parts.append(last)
        formatted = " ".join(part for part in parts if part)
        return keep_richer_spelling(formatted, current) if formatted else current

    if last and token_set(last).issubset(pref_tokens):
        core = preferred
    elif last:
        core = f"{preferred} {last}"
    else:
        core = preferred

    extra = leftover_given(first, middle, preferred)
    if extra:
        return f"{core} ({extra})"
    return core


def load_roster_rows(xlsx: Path) -> list[dict]:
    workbook = load_workbook(xlsx, data_only=True)
    rows: list[dict] = []
    try:
        ib2 = workbook["IB2s"]
        for values in ib2.iter_rows(min_row=2, values_only=True):
            first = clean_part(values[1] if len(values) > 1 else "")
            middle = clean_part(values[2] if len(values) > 2 else "")
            last = clean_part(values[3] if len(values) > 3 else "")
            preferred = clean_part(values[4] if len(values) > 4 else "")
            if not any((first, middle, last, preferred)):
                continue
            rows.append(
                {
                    "cohort": "IB2",
                    "first": first,
                    "middle": middle,
                    "last": last,
                    "preferred": preferred,
                    "country": None,
                    "legal": " ".join(part for part in (first, middle, last) if part),
                }
            )

        ib1 = workbook["IB1s"]
        for values in ib1.iter_rows(min_row=2, values_only=True):
            first = clean_part(values[0] if len(values) > 0 else "")
            last = clean_part(values[1] if len(values) > 1 else "")
            preferred = clean_part(values[2] if len(values) > 2 else "")
            country = country_from_nationality(values[3] if len(values) > 3 else None)
            if not any((first, last, preferred)):
                continue
            rows.append(
                {
                    "cohort": "IB1",
                    "first": first,
                    "middle": "",
                    "last": last,
                    "preferred": preferred,
                    "country": country,
                    "legal": " ".join(part for part in (first, last) if part),
                }
            )
    finally:
        workbook.close()
    return rows


def score_match(row: dict, student: dict) -> float:
    if row["cohort"] != student["cohort"]:
        return 0.0
    student_tokens = token_set(strip_parens(student["name"]))
    legal_tokens = token_set(row["legal"])
    pref_tokens = token_set(row["preferred"])
    last_tokens = token_set(row["last"])
    if not student_tokens:
        return 0.0

    union = legal_tokens | student_tokens | pref_tokens
    overlap = student_tokens & (legal_tokens | pref_tokens)
    if not union:
        return 0.0
    jaccard = len(overlap) / len(union)

    last_hit = 0.25 if last_tokens and last_tokens & student_tokens else 0.0
    pref_hit = 0.2 if pref_tokens and pref_tokens & student_tokens else 0.0
    subset_bonus = 0.0
    if legal_tokens and (
        legal_tokens.issubset(student_tokens) or student_tokens.issubset(legal_tokens)
    ):
        subset_bonus = 0.35
    if pref_tokens and pref_tokens.issubset(student_tokens):
        subset_bonus = max(subset_bonus, 0.2)

    # Require some real overlap so "Sofi" rows don't steal random Sofias.
    if len(overlap) < 1:
        return 0.0
    if last_tokens and not (last_tokens & student_tokens) and len(overlap) < 2:
        return 0.0
    return jaccard + last_hit + pref_hit + subset_bonus


def match_roster(rows: list[dict], students: list[dict]) -> tuple[dict[str, dict], list[dict], list[dict]]:
    remaining = {student["id"]: student for student in students}
    ranked: list[tuple[float, int, str, dict]] = []
    for index, row in enumerate(rows):
        for student in students:
            value = score_match(row, student)
            if value > 0.55:
                ranked.append((value, index, student["id"], row))
    ranked.sort(key=lambda item: (-item[0], item[1]))

    assigned: dict[str, dict] = {}
    used_rows: set[int] = set()
    for value, index, student_id, row in ranked:
        if student_id in assigned or index in used_rows or student_id not in remaining:
            continue
        assigned[student_id] = row
        used_rows.add(index)
        remaining.pop(student_id)

    unmatched_rows = [row for i, row in enumerate(rows) if i not in used_rows]
    unmatched_students = list(remaining.values())
    return assigned, unmatched_rows, unmatched_students


def apply_preferred_names(students: list[dict], xlsx: Path = ROSTER_XLSX) -> dict:
    rows = load_roster_rows(xlsx)
    assigned, unmatched_rows, unmatched_students = match_roster(rows, students)
    changes: list[dict] = []
    updated: list[dict] = []
    for student in students:
        row = assigned.get(student["id"])
        if row:
            name = format_display(
                row["preferred"],
                row["first"],
                row["middle"],
                row["last"],
                student["name"],
            )
            country = row["country"]
        else:
            name = student["name"]
            country = student.get("country") or None
        if name != student["name"] or country != student.get("country"):
            changes.append(
                {
                    "id": student["id"],
                    "from": student["name"],
                    "to": name,
                    "country": country,
                }
            )
        next_student = {
            "id": student["id"],
            "name": name,
            "cohort": student["cohort"],
            "country": country,
            "house": student.get("house"),
            "blocks": student["blocks"],
        }
        for key, value in student.items():
            if key not in next_student:
                next_student[key] = value
        updated.append(next_student)
    return {
        "students": updated,
        "changes": changes,
        "unmatched_rows": unmatched_rows,
        "unmatched_students": unmatched_students,
        "matched": len(assigned),
    }


def split_jammed_name(name: str) -> str:
    name = re.sub(r"([a-z])([A-Z])", r"\1 \2", name)
    pieces: list[str] = []
    for word in name.split():
        lower = word.lower()
        split = False
        for particle in ("del", "de", "van", "von"):
            if (
                lower != particle
                and lower.endswith(particle)
                and len(lower) - len(particle) >= 3
            ):
                pieces.append(word[: -len(particle)])
                pieces.append(word[-len(particle) :])
                split = True
                break
        if not split:
            pieces.append(word)
    return " ".join(pieces)


def normalize_house(value: object) -> str | None:
    text = clean_part(value)
    if not text:
        return None
    key = strip_accents(text).lower()
    return HOUSE_NAMES.get(key, smart_case(text))


def year_to_cohort(value: object) -> str | None:
    text = clean_part(value).upper().replace(" ", "")
    if "IB1" in text or text.endswith("-1") or text == "1":
        return "IB1"
    if "IB2" in text or text.endswith("-2") or text == "2":
        return "IB2"
    return None


def load_house_rows(xlsx: Path) -> list[dict]:
    workbook = load_workbook(xlsx, data_only=True)
    rows: list[dict] = []
    try:
        for sheet in workbook.worksheets:
            for values in sheet.iter_rows(min_row=2, values_only=True):
                raw_name = split_jammed_name(clean_part(values[0] if values else ""))
                if not raw_name or raw_name.lower() == "name":
                    continue
                cohort = year_to_cohort(values[1] if len(values) > 1 else None)
                house = normalize_house(values[2] if len(values) > 2 else None)
                country = clean_part(values[3] if len(values) > 3 else None) or None
                rows.append(
                    {
                        "name": raw_name,
                        "cohort": cohort,
                        "house": house,
                        "country": country,
                        "sheet": sheet.title,
                    }
                )
    finally:
        workbook.close()
    return rows


def score_house_match(row: dict, student: dict) -> float:
    if row["cohort"] and row["cohort"] != student["cohort"]:
        return 0.0
    row_tokens = token_set(row["name"])
    student_tokens = token_set(student["name"])
    if not row_tokens or not student_tokens:
        return 0.0
    overlap = row_tokens & student_tokens
    if not overlap:
        return 0.0
    union = row_tokens | student_tokens
    jaccard = len(overlap) / len(union)
    subset = (
        0.4
        if row_tokens.issubset(student_tokens) or student_tokens.issubset(row_tokens)
        else 0.0
    )
    last_row = tokens(strip_parens(row["name"]))[-1:]
    last_student = tokens(strip_parens(student["name"]))[-1:]
    last_hit = 0.3 if last_row and last_row == last_student else 0.0
    given = set(tokens(row["name"])[:2])
    given_hit = 0.25 if len(given) >= 2 and given.issubset(student_tokens) else 0.0
    if len(overlap) < 2 and not last_hit and not subset:
        return 0.0
    return jaccard + subset + last_hit + given_hit


def apply_houses(students: list[dict], xlsx: Path = HOUSES_XLSX) -> dict:
    rows = load_house_rows(xlsx)
    remaining = {student["id"]: student for student in students}
    ranked: list[tuple[float, int, str, dict]] = []
    for index, row in enumerate(rows):
        for student in students:
            value = score_house_match(row, student)
            if value > 0.45:
                ranked.append((value, index, student["id"], row))
    ranked.sort(key=lambda item: (-item[0], item[1]))

    assigned: dict[str, dict] = {}
    used_rows: set[int] = set()
    for _value, index, student_id, row in ranked:
        if student_id in assigned or index in used_rows or student_id not in remaining:
            continue
        assigned[student_id] = row
        used_rows.add(index)
        remaining.pop(student_id)

    updated: list[dict] = []
    for student in students:
        row = assigned.get(student["id"])
        house = row["house"] if row else student.get("house")
        next_student = {
            "id": student["id"],
            "name": student["name"],
            "cohort": student["cohort"],
            "country": student.get("country"),
            "house": house,
            "blocks": student["blocks"],
        }
        for key, value in student.items():
            if key not in next_student:
                next_student[key] = value
        updated.append(next_student)

    unmatched_rows = [row for i, row in enumerate(rows) if i not in used_rows]
    unmatched_students = list(remaining.values())
    return {
        "students": updated,
        "matched": len(assigned),
        "assigned": assigned,
        "unmatched_rows": unmatched_rows,
        "unmatched_students": unmatched_students,
    }


def parse_student_emails(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    emails: dict[str, str] = {}
    for match in re.finditer(r'"([^"]+)":\s*"([^"]+@[^"]+)"', text):
        emails[match.group(1)] = match.group(2)
    return emails


def write_edge_roster(students: list[dict], emails: dict[str, str], path: Path) -> None:
    entries = []
    for student in students:
        email = emails.get(student["id"])
        email_js = json.dumps(email) if email else "null"
        entries.append(
            "  {\n"
            f'    "id": {json.dumps(student["id"])},\n'
            f'    "name": {json.dumps(student["name"], ensure_ascii=False)},\n'
            f'    "cohort": {json.dumps(student["cohort"])},\n'
            f'    "email": {email_js}\n'
            "  }"
        )
    body = ",\n".join(entries)
    path.write_text(
        "export type RosterStudent = {\n"
        "  id: string;\n"
        "  name: string;\n"
        '  cohort: "IB1" | "IB2";\n'
        "  email: string | null;\n"
        "};\n"
        "\n"
        "export const ROSTER: RosterStudent[] = [\n"
        f"{body}\n"
        "];\n"
        "\n"
        "export function rosterByEmail(email: string): RosterStudent | null {\n"
        "  const lower = email.trim().toLowerCase();\n"
        "  if (!lower) return null;\n"
        "  return ROSTER.find((student) => student.email?.toLowerCase() === lower) ?? null;\n"
        "}\n",
        encoding="utf-8",
    )


def apply_to_payload(payload: dict, xlsx: Path = ROSTER_XLSX) -> dict:
    result = apply_preferred_names(payload["students"], xlsx)
    houses = apply_houses(result["students"])
    payload = dict(payload)
    payload["students"] = houses["students"]
    return payload


def main() -> None:
    payload = json.loads(STUDENTS_JSON.read_text(encoding="utf-8"))
    result = apply_preferred_names(payload["students"])
    houses = apply_houses(result["students"])
    payload["students"] = houses["students"]
    source = str(payload.get("source") or "")
    extras = [ROSTER_XLSX.name, HOUSES_XLSX.name]
    for name in extras:
        if name not in source:
            source = f"{source} + {name}" if source else name
    payload["source"] = source
    STUDENTS_JSON.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    emails = parse_student_emails(EMAILS_TS)
    write_edge_roster(payload["students"], emails, EDGE_ROSTER)

    print(f"Matched {result['matched']} students")
    print(f"Name/country changes: {len(result['changes'])}")
    print(f"Unmatched roster rows: {len(result['unmatched_rows'])}")
    for row in result["unmatched_rows"]:
        print(f"  roster {row['cohort']}: {row['preferred'] or row['first']} | {row['legal']}")
    print(f"Unmatched students: {len(result['unmatched_students'])}")
    for student in result["unmatched_students"]:
        print(f"  {student['id']}: {student['name']}")
    print(f"Houses matched {houses['matched']}")
    print(f"House roster skipped: {len(houses['unmatched_rows'])}")
    for row in houses["unmatched_rows"]:
        print(f"  {row['name']} ({row['house']})")
    print(f"House student skipped: {len(houses['unmatched_students'])}")
    for student in houses["unmatched_students"]:
        print(f"  {student['id']}: {student['name']}")
    print("\nChanges:")
    for item in result["changes"]:
        mark = "" if item["from"] == item["to"] else f"{item['from']} -> {item['to']}"
        if item["from"] == item["to"]:
            mark = f"{item['to']} (country {item['country']})"
        print(f"  {item['id']}: {mark}")


if __name__ == "__main__":
    main()
