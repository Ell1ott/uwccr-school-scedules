#!/usr/bin/env python3
"""Add a Full Name column to IB2 list.xlsx."""

from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(__file__).resolve().parents[1] / "IB2 list.xlsx"


def part(value) -> str:
    text = "" if value is None else str(value).strip()
    return "" if text.upper() in {"", "X"} else text


wb = load_workbook(XLSX)
ws = wb.active
ws["F1"] = "Full Name"

for row in range(2, ws.max_row + 1):
    names = [part(ws.cell(row, col).value) for col in (2, 3, 4)]
    full = " ".join(n for n in names if n)
    preferred = part(ws.cell(row, 5).value)
    ws.cell(row, 6).value = f"{full} ({preferred})" if preferred else full

wb.save(XLSX)
print(f"Wrote Full Name column to {XLSX.name}")
